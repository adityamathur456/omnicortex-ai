from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE_COLOR = RGBColor(15, 23, 42)
ACCENT_COLOR = RGBColor(13, 148, 136)
MUTED_COLOR = RGBColor(100, 116, 139)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = TITLE_COLOR


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = TITLE_COLOR
    return paragraph


def add_section(document, title):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR
    return paragraph


def add_key_value_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    header = table.rows[0].cells
    header[0].text = "Area"
    header[1].text = "Details"
    for cell in header:
      shade_cell(cell, "E2F7F5")
      set_cell_margins(cell)
      cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
      for paragraph in cell.paragraphs:
          for run in paragraph.runs:
              run.font.bold = True
              run.font.color.rgb = TITLE_COLOR

    for left, right in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = left
        row_cells[1].text = right
        for cell in row_cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = TITLE_COLOR
    return table


document = Document()

section = document.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = document.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("Omnicortex AI Frontend Documentation")
title_run.bold = True
title_run.font.size = Pt(22)
title_run.font.color.rgb = TITLE_COLOR

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run(
    "Architecture, user flows, API integration, history management, and runtime behavior"
)
subtitle_run.italic = True
subtitle_run.font.size = Pt(11)
subtitle_run.font.color.rgb = MUTED_COLOR

meta = document.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_run = meta.add_run(
    "Project folder: C:\\Users\\adity\\OneDrive\\Documents\\Omnicortex-Client"
)
meta_run.font.size = Pt(9.5)
meta_run.font.color.rgb = ACCENT_COLOR

document.add_paragraph("")

add_section(document, "1. Frontend Overview")
add_body(
    document,
    "Omnicortex AI is a React + Vite single-page application that provides three authenticated AI workflows inside one SaaS dashboard: resume generation, image generation, and medical analysis. The application uses TypeScript for typing discipline, Tailwind CSS for interface styling, Axios for API access, React Router for page navigation, and localStorage-backed state for login persistence and generated link history.",
)
for item in [
    "Public routes: /login and /register",
    "Protected routes: /dashboard, /resume, /image, and /medical",
    "Primary layout: sidebar navigation, top user bar, responsive content area",
    "Frontend brand: Omnicortex AI with SVG logo and favicon integration",
]:
    add_bullet(document, item)

add_section(document, "2. Technology Stack")
add_key_value_table(
    document,
    [
        ("Framework", "React 19 with Vite 6"),
        ("Language", "TypeScript"),
        ("Styling", "Tailwind CSS"),
        ("Routing", "React Router DOM"),
        ("HTTP client", "Axios with request interceptor"),
        ("Notifications", "react-hot-toast"),
        ("Icons", "lucide-react"),
        ("Preview behavior", "Iframe for resume PDF preview and inline img preview for images"),
    ],
)

add_section(document, "3. Directory and Code Organization")
add_body(
    document,
    "The frontend is intentionally separated by responsibility so API access, layout components, page screens, persistence logic, and low-level utilities remain isolated.",
)
structure = document.add_paragraph()
structure.paragraph_format.space_after = Pt(8)
structure_run = structure.add_run(
    "src/\n"
    "  api/          backend integration modules\n"
    "  components/   shared UI, layout, route guards, branding\n"
    "  pages/        route-level screens\n"
    "  store/        auth context and generation-history persistence\n"
    "  utils/        download helpers, constants, storage, error parsing\n"
)
structure_run.font.name = "Courier New"
structure_run.font.size = Pt(9.5)
structure_run.font.color.rgb = TITLE_COLOR

add_section(document, "4. Authentication and Session Flow")
add_body(
    document,
    "Authentication is managed through a React context in src/store/auth-store.tsx. On login, the frontend posts credentials to /auth/login, normalizes the response into an AuthSession object, stores the access token plus user details in localStorage, and exposes the session to the rest of the app through useAuth(). On logout, the session state and storage entry are cleared.",
)
add_key_value_table(
    document,
    [
        ("Register API", "POST http://127.0.0.1:8000/auth/register"),
        ("Login API", "POST http://127.0.0.1:8000/auth/login"),
        ("Stored auth key", "omnicortex-ai-auth"),
        ("Token usage", "Axios interceptor adds Authorization: Bearer <token> to requests"),
        ("Route guard behavior", "Unauthenticated access to protected pages redirects to /login"),
    ],
)

add_section(document, "5. Routing Model")
add_body(
    document,
    "Routing is declared centrally in src/App.tsx. PublicRoute prevents authenticated users from returning to login/register, while ProtectedRoute prevents unauthenticated users from entering dashboard tools.",
)
add_key_value_table(
    document,
    [
        ("Public", "/login, /register"),
        ("Protected", "/dashboard, /resume, /image, /medical"),
        ("Fallback", "/ redirects to /dashboard, unmatched paths render 404 screen"),
    ],
)

add_section(document, "6. Feature Pages")
add_body(
    document,
    "Each feature page follows the same operating model: form inputs on the left, results on the right, loading states through skeleton placeholders, toast notifications for success/error, and API integration handled by dedicated modules under src/api.",
)
for item in [
    "Resume Generator: sends raw text/plain body to /generate-resume, previews returned pdf_url in an iframe, stores PDF history, and allows open-in-new-tab plus download actions.",
    "Image Generator: sends prompt/mode/steps JSON to /generate-image, uses backend-supported mode values (simple, portrait, widescreen_landscape, product_photo), previews the image, and stores image history.",
    "Medical Analyzer: submits multipart form data with file and modality_hint to /analyze-medical, then renders a structured result card with preview image, findings, conditions, recommendation, and disclaimer.",
]:
    add_bullet(document, item)

add_section(document, "7. URL History and Persistence")
add_body(
    document,
    "The frontend preserves generated links and page state in localStorage through src/store/generation-history.ts. This was added so navigation between Resume and Image pages does not discard previously generated outputs. Each page restores its last-known form values and most recent result on load, and it also maintains a clickable history list.",
)
add_key_value_table(
    document,
    [
        ("Resume state key", "omnicortex-ai-resume-page"),
        ("Image state key", "omnicortex-ai-image-page"),
        ("History size cap", "8 items per tool"),
        ("Resume history fields", "id, createdAt, input, pdfUrl, result"),
        ("Image history fields", "id, createdAt, prompt, mode, steps, imageUrl"),
    ],
)
add_body(
    document,
    "Resume and image pages call getResumePageState()/getImagePageState() during initialization, then use useEffect() to keep the current page state synchronized back into localStorage whenever inputs, result URLs, or history arrays change. History appenders also deduplicate by URL before trimming to the maximum size.",
)

add_section(document, "8. Download Management")
add_body(
    document,
    "The current download implementation uses a same-origin proxy path instead of direct cross-origin file saving. This is defined in src/utils/download.ts and supported by a Vite middleware in vite.config.ts. The frontend rewrites download requests to /__download?url=<remote-url>&filename=<name>, and the Vite server fetches the remote file, applies a safe filename, and returns it with Content-Disposition set for browser download.",
)
for item in [
    "Reason for proxy: remote object storage URLs may preview correctly but still fail direct browser download due to cross-origin restrictions or missing Content-Disposition.",
    "Client helper: buildDownloadProxyUrl() generates the same-origin route.",
    "Client trigger: downloadFromUrl() creates an anchor against the proxy URL so the browser treats it as a first-party file download.",
    "Filename handling: getFilenameFromUrl() extracts a readable default filename from the returned URL path.",
]:
    add_bullet(document, item)

add_section(document, "9. User Interface and Styling")
add_body(
    document,
    "The UI is designed as a restrained SaaS dashboard rather than a marketing landing page. Tailwind utility classes define spacing, borders, typography, and responsive behavior. Shared components include Button, Card, Input, Select, Textarea, Skeleton, Badge, PageHeader, EmptyState, AuthShell, and AppShell. This keeps the styling consistent across all workflows and reduces duplicated page-specific markup.",
)

add_section(document, "10. Error Handling and Feedback")
add_body(
    document,
    "All API pages use toast notifications for visible status feedback. The helper in src/utils/error.ts normalizes backend and Axios errors into short user-facing strings. Loading conditions are surfaced through button spinners and skeleton placeholders so the user can distinguish waiting states from broken states.",
)

add_section(document, "11. How to Run and Build")
add_key_value_table(
    document,
    [
        ("Development", "npm run dev or corepack pnpm dev"),
        ("Type check", "tsc -b"),
        ("Production build", "vite build"),
        ("Server host", "127.0.0.1:5173"),
        ("Backend base URL", "http://127.0.0.1:8000"),
    ],
)

add_section(document, "12. Practical Maintenance Notes")
for item in [
    "If backend auth response changes, login normalization logic in src/api/auth.ts is the adjustment point.",
    "If new tools are added, add a page, API module, sidebar item, and protected route entry.",
    "If more persisted tools are introduced, follow the same localStorage pattern used in src/store/generation-history.ts.",
    "If download behavior changes again, check both src/utils/download.ts and the /__download middleware in vite.config.ts.",
]:
    add_bullet(document, item)

document.add_paragraph("")
closing = document.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
closing_run = closing.add_run("Prepared for Omnicortex AI frontend handoff and maintenance.")
closing_run.font.size = Pt(10)
closing_run.font.italic = True
closing_run.font.color.rgb = MUTED_COLOR

output_path = r"C:\Users\adity\OneDrive\Documents\Omnicortex-Client\Omnicortex-Frontend-Documentation.docx"
document.save(output_path)
print(output_path)
